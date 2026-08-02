"""Shared sample veterinary record content for PDF and DOCX test fixtures."""

from __future__ import annotations

from io import BytesIO

SAMPLE_VET_RECORD_TEXT = """Sunshine Vet Clinic
Visit date: 2024-06-10
Veterinarian: Dr. Smith
Pet: Buddy, Canine, Labrador Retriever, Male
Date of birth: 2020-03-15
Owner: Jane Doe, +1-555-0100, jane@example.com
Chief complaint: Left ear scratching and head shaking
History: Symptoms for 3 days
Examination: Mild erythema in left ear canal
Diagnosis: Otitis externa
Treatment: Topical ear medication
Medication: Otomax, 4 drops, Twice daily for 7 days
Notes: Follow up in 1 week
"""


def make_sample_pdf_bytes() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in SAMPLE_VET_RECORD_TEXT.splitlines():
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def make_sample_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    for line in SAMPLE_VET_RECORD_TEXT.splitlines():
        doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
